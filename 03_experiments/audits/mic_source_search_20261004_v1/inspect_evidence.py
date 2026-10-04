"""Read original source files, retain unresolved evidence, and issue no labels."""
import csv
import datetime
import hashlib
import json
import math
from pathlib import Path
import re
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / 'inspection_v2'
PROJECT = ROOT.parents[2]
FORMAL = PROJECT / '03_experiments/prepared/amd02_evidence_clean_20261001/records.csv'
MIC = PROJECT / '03_experiments/audits/mic_followup_release_20261004_v1/evidence/source_review.csv'

def text(node):
    return ' '.join(node.itertext()) if node is not None else ''

def write(name, value):
    with (OUTPUT / name).open('x', encoding='utf-8') as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2, allow_nan=False)

def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def main():
    outputs = ['source_inventory.json', 'archaeome_observations.json', 'inspection.json']
    if OUTPUT.exists():
        raise SystemExit('Refusing to overwrite inspection outputs')
    OUTPUT.mkdir()
    with FORMAL.open(encoding='utf-8-sig', newline='') as stream:
        development = list(csv.DictReader(stream))
    with MIC.open(encoding='utf-8-sig', newline='') as stream:
        mic_sources = list(csv.DictReader(stream))
    decisions = {
        'archaeome2025': ('priority_primary_measurement_review',
            '80 experimental IDs; sequence-to-ID certificates not obtained; termini unknown; '
            'author HC50 fits outside the tested range cannot be accepted as exact measurements'),
        'foundation2025': ('priority_primary_measurement_review',
            'Table 3 has 18 E. coli MIC values; AI18 and KW13 HC50 claims require Fig S4; '
            'sequence, chemical form, supplements and all-member isolation remain unverified'),
        'pandi2023': ('not_independent_formal_external_source',
            'PMID37938588 already represented in the formal development set; '
            'synthetic assay peptides use TentaGel S RAM resin, requiring amidation audit'),
        'hybrid2018': ('human_hc50_ineligible_mic_only_pending',
            'Table 1 and Methods 4.8 explicitly use sheep erythrocytes; '
            'MIC evidence may be reviewed separately; chemical identity unresolved'),
        'hydramp2023': ('human_hc50_ineligible_mic_only_pending',
            'Methods Sec35 use sheep RBCs; Sec33 uses Rink Amide resin; '
            'cannot transfer to human HC50 or unmodified MIC molecule domain'),
        'capitellacin2022': ('molecule_domain_barrier',
            'Disulfide beta-hairpin peptide family; Fig2 hemolysis and Methods timing differ; '
            'not a verified unmodified linear cohort'),
        'spectrum2020': ('hc50_extrapolation_barrier',
            'Methods 2.5 explicitly describe extrapolated HC50; no automatic release of fitted '
            'values as directly observed endpoints; chemical identity and isolation unresolved'),
        'qmap2026': ('secondary_discovery_resource_only',
            'DBAASP-derived aggregation, including nonhuman RBC fallback and averaged endpoints; '
            'must return to underlying primary observations; not an independent measurement study'),
    }
    inventory = []
    for path in sorted((ROOT / 'raw').glob('*.xml')):
        article = ET.parse(path).getroot()
        identifiers = {node.get('pub-id-type'): node.text
                       for node in article.findall('.//article-meta/article-id')}
        pmid, doi = identifiers.get('pmid'), identifiers.get('doi')
        matches = []
        for row in development:
            serialized = json.dumps(row, ensure_ascii=False).lower()
            if ((doi and doi.lower() in serialized) or
                    (pmid and re.search(r'(?<!\d)' + re.escape(pmid) + r'(?!\d)', serialized))):
                matches.append(row.get('record_id') or row.get('id') or next(iter(row.values())))
        mic_matches = [row['source_id'] for row in mic_sources
                       if doi and doi.lower() in json.dumps(row).lower()]
        disposition, barrier = decisions[path.stem]
        inventory.append({'source_id': path.stem, 'doi': doi, 'pmid': pmid,
                          'pmcid': identifiers.get('pmcid'),
                          'title': text(article.find('.//article-title')),
                          'primary_url': 'https://doi.org/' + doi,
                          'license_text': text(article.find('.//license')),
                          'primary_file': str(path.relative_to(ROOT)), 'sha256': digest(path),
                          'formal_development_study_matches': matches,
                          'mic_development_study_matches': mic_matches,
                          'study_check_scope': 'Exact DOI/PMID occurrence; full alias audit pending',
                          'disposition': disposition, 'barrier': barrier,
                          'approved_for_training': False,
                          'all_member_homology_status': 'not_run_missing_confirmed_molecular_identity'})
    write('source_inventory.json', inventory)

    mic_path = ROOT / 'raw/41564_2025_2061_MOESM8_ESM.xlsx'
    hc_path = ROOT / 'raw/41564_2025_2061_MOESM10_ESM.xlsx'
    mic_book = openpyxl.load_workbook(mic_path, read_only=True, data_only=False)
    mic_sheet = mic_book['Fig 2a']
    rows = list(mic_sheet.iter_rows(values_only=True))
    assert len(rows) == 81 and len(rows[0]) == 12
    assert len({row[0] for row in rows[1:]}) == 80
    hc_book = openpyxl.load_workbook(hc_path, read_only=True, data_only=False)
    hc_sheet = hc_book['Fig 4c']
    hc_rows = list(hc_sheet.iter_rows(values_only=True))
    assert len(hc_rows) == 2 and len(hc_rows[0]) == 81
    assert hc_rows[1][0] == 'HC50'
    # Normalize the author's spelling only for a unique ID join, never row-order guessing.
    hc_lookup = {}
    for col in range(2, 82):
        raw_id = hc_rows[0][col - 1]
        normalized = raw_id.replace('Archeasin-', 'Archaeasin-')
        assert normalized not in hc_lookup
        hc_lookup[normalized] = {'original_id': raw_id,
                                'value': hc_rows[1][col - 1],
                                'cell': get_column_letter(col) + '2'}
    assert set(hc_lookup) == {row[0] for row in rows[1:]}
    observations = []
    for row_index, row in enumerate(rows[1:], start=2):
        name = row[0]
        for column in (3, 4, 5):
            value = row[column - 1]
            observations.append({'experimental_id': name, 'endpoint': 'MIC',
                'target': rows[0][column - 1], 'raw_value': value, 'unit': 'uM',
                'raw_cell': get_column_letter(column) + str(row_index),
                'sheet': 'Fig 2a', 'source_file': str(mic_path.relative_to(ROOT)),
                'interpretation': 'numeric_original_endpoint' if value is not None else
                    'blank_in_source_no_value_imputed',
                'sequence': None, 'chemical_form': None, 'derived_label': None,
                'admission': 'pending_molecule_identity_and_isolation'})
        item = hc_lookup[name]
        raw = item['value']
        numeric = float(raw) if raw is not None else None
        observations.append({'experimental_id': name, 'original_id': item['original_id'],
            'endpoint': 'HC50_author_linear_regression_output',
            'target': 'human erythrocytes', 'raw_value': raw, 'unit': 'uM',
            'raw_cell': item['cell'], 'sheet': 'Fig 4c',
            'source_file': str(hc_path.relative_to(ROOT)),
            'author_fit_exceeds_64_uM_figure_range': numeric > 64 if numeric is not None else None,
            'implausible_fit_above_1e6_uM': numeric > 1e6 if numeric is not None else None,
            'sequence': None, 'chemical_form': None, 'derived_label': None,
            'admission': 'not_released_author_fit_requires_dose_response_review',
            'conflict': 'Workbook Fig4c HC50 vs main figure caption panel4d; '
                        'methods concentration range differs from Extended Data Fig7 caption'})
    assert len(observations) == 320
    write('archaeome_observations.json', observations)
    rejected_asset = ROOT / 'raw/41564_2025_2061_MOESM17_ESM.xlsx'
    assert rejected_asset.read_bytes().lstrip().startswith(b'<!DOCTYPE html>')
    result = {'generated_at': datetime.datetime.now(datetime.timezone.utc).isoformat(),
              'status': 'source_discovery_completed_evidence_release_pending',
              'review_mode': 'agent_primary_source_review_not_independent_human_review',
              'primary_articles_archived': len(inventory),
              'formal_development_records_checked': len(development),
              'archaeome_experimental_ids': 80,
              'ecoli_mic_slots': 240,
              'ecoli_numeric_original_values': sum(x['endpoint'] == 'MIC' and x['raw_value'] is not None for x in observations),
              'ecoli_blank_slots': sum(x['endpoint'] == 'MIC' and x['raw_value'] is None for x in observations),
              'human_hc50_author_fit_outputs': 80,
              'hc50_author_fits_exceed_figure_range': sum(bool(x.get('author_fit_exceeds_64_uM_figure_range')) for x in observations),
              'implausible_hc50_fits': sum(bool(x.get('implausible_fit_above_1e6_uM')) for x in observations),
              'prediction_only_csv_rows_not_released': 265,
              'approved_new_molecules': 0,
              'formal_external_cohort_released': False,
              'certificate_repository': 'https://data.mendeley.com/datasets/d8yzgtdrcp/3',
              'certificate_retrieval': 'HTTP403 on public page and file-list API; files not obtained',
              'html_masquerading_as_xlsx_rejected': str(rejected_asset.relative_to(ROOT)),
              'inputs': {str(FORMAL.relative_to(PROJECT)): digest(FORMAL),
                         str(MIC.relative_to(PROJECT)): digest(MIC)},
              'artifacts': {name: digest(OUTPUT / name) for name in outputs[:2]}}
    write('inspection.json', result)
    print(json.dumps(result, ensure_ascii=False, indent=2))

if __name__ == '__main__':
    main()
